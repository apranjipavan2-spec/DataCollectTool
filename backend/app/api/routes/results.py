"""Results Framework: logframe (Goal->Outcome->Output->Activity), indicators, and the
Indicator Tracking Table (ITT) — baseline/midline/endline actual-vs-target values."""
import io
from datetime import datetime, timezone
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.deps import require_enumerator, require_supervisor
from app.core.soft_delete import soft_delete
from app.models.program import Program, ProgramQuestionnaire
from app.models.results_framework import LogframeLevel, Indicator, IndicatorValue
from app.models.submission import Submission
from app.services.indicator_compute import compute_indicator, rollup_tree

router = APIRouter()


# ── Schemas ──────────────────────────────────────────────────────────────────

class LogframeLevelIn(BaseModel):
    parent_id: Optional[str] = None
    level_type: str  # goal|outcome|output|activity
    title: str
    description: str = ""
    sort_order: int = 0


class IndicatorIn(BaseModel):
    logframe_level_id: str
    code: str = ""
    name: str
    unit: str = "count"          # count|percent|ratio|mean|text
    direction: str = "increase"  # increase|decrease
    disaggregate_by: list[str] = []
    source_form_field: Optional[str] = None
    aggregation: str = "count"   # count|sum|mean|percent
    numerator_filter: Optional[dict] = None
    denominator_filter: Optional[dict] = None
    value_source: str = "auto"   # auto|manual


class IndicatorValueOverrideIn(BaseModel):
    target_value: Optional[float] = None
    actual_value: Optional[float] = None
    note: Optional[str] = None


def _get_program(prog_id: str, user: dict, db: Session) -> Program:
    p = db.query(Program).filter(Program.id == prog_id, Program.tenant_id == user["tenant_id"]).first()
    if not p:
        raise HTTPException(404, "Program not found")
    return p


# ── Logframe CRUD ────────────────────────────────────────────────────────────

@router.get("/programs/{prog_id}/logframe")
def get_logframe(prog_id: str, user=Depends(require_enumerator), db: Session = Depends(get_db)):
    _get_program(prog_id, user, db)
    levels = db.query(LogframeLevel).filter(
        LogframeLevel.program_id == prog_id, LogframeLevel.tenant_id == user["tenant_id"],
    ).order_by(LogframeLevel.sort_order).all()
    return [{"id": str(l.id), "parent_id": str(l.parent_id) if l.parent_id else None,
             "level_type": l.level_type, "title": l.title, "description": l.description,
             "sort_order": l.sort_order} for l in levels]


@router.post("/programs/{prog_id}/logframe")
def create_logframe_level(prog_id: str, body: LogframeLevelIn, user=Depends(require_supervisor), db: Session = Depends(get_db)):
    _get_program(prog_id, user, db)
    lvl = LogframeLevel(tenant_id=user["tenant_id"], program_id=prog_id, **body.model_dump())
    db.add(lvl); db.commit(); db.refresh(lvl)
    return {"id": str(lvl.id)}


@router.put("/programs/{prog_id}/logframe/{level_id}")
def update_logframe_level(prog_id: str, level_id: str, body: LogframeLevelIn, user=Depends(require_supervisor), db: Session = Depends(get_db)):
    lvl = db.query(LogframeLevel).filter(LogframeLevel.id == level_id, LogframeLevel.program_id == prog_id,
                                          LogframeLevel.tenant_id == user["tenant_id"]).first()
    if not lvl: raise HTTPException(404, "Logframe level not found")
    for k, v in body.model_dump().items():
        setattr(lvl, k, v)
    db.commit()
    return {"id": str(lvl.id)}


@router.delete("/programs/{prog_id}/logframe/{level_id}")
def delete_logframe_level(prog_id: str, level_id: str, user=Depends(require_supervisor), db: Session = Depends(get_db)):
    lvl = db.query(LogframeLevel).filter(LogframeLevel.id == level_id, LogframeLevel.program_id == prog_id,
                                          LogframeLevel.tenant_id == user["tenant_id"]).first()
    if not lvl: raise HTTPException(404, "Logframe level not found")
    soft_delete(lvl); db.commit()
    return {"deleted": True}


# ── Indicator CRUD ───────────────────────────────────────────────────────────

@router.get("/programs/{prog_id}/indicators")
def list_indicators(prog_id: str, user=Depends(require_enumerator), db: Session = Depends(get_db)):
    _get_program(prog_id, user, db)
    inds = db.query(Indicator).filter(Indicator.program_id == prog_id, Indicator.tenant_id == user["tenant_id"]).all()
    return [{"id": str(i.id), "logframe_level_id": str(i.logframe_level_id), "code": i.code,
             "name": i.name, "unit": i.unit, "direction": i.direction,
             "disaggregate_by": i.disaggregate_by, "source_form_field": i.source_form_field,
             "aggregation": i.aggregation, "numerator_filter": i.numerator_filter,
             "denominator_filter": i.denominator_filter, "value_source": i.value_source} for i in inds]


@router.post("/programs/{prog_id}/indicators")
def create_indicator(prog_id: str, body: IndicatorIn, user=Depends(require_supervisor), db: Session = Depends(get_db)):
    _get_program(prog_id, user, db)
    ind = Indicator(tenant_id=user["tenant_id"], program_id=prog_id, **body.model_dump())
    db.add(ind); db.flush()
    # Seed one blank ITT row (disaggregation={}) per existing wave so target/actual are
    # editable immediately — auto-compute fills actual_value later for auto indicators.
    waves = db.query(ProgramQuestionnaire).filter(
        ProgramQuestionnaire.program_id == prog_id, ProgramQuestionnaire.tenant_id == user["tenant_id"]).all()
    for w in waves:
        db.add(IndicatorValue(tenant_id=user["tenant_id"], indicator_id=ind.id, questionnaire_id=w.id,
                               disaggregation={}, value_source=ind.value_source))
    db.commit(); db.refresh(ind)
    return {"id": str(ind.id)}


@router.put("/programs/{prog_id}/indicators/{ind_id}")
def update_indicator(prog_id: str, ind_id: str, body: IndicatorIn, user=Depends(require_supervisor), db: Session = Depends(get_db)):
    ind = db.query(Indicator).filter(Indicator.id == ind_id, Indicator.program_id == prog_id,
                                      Indicator.tenant_id == user["tenant_id"]).first()
    if not ind: raise HTTPException(404, "Indicator not found")
    for k, v in body.model_dump().items():
        setattr(ind, k, v)
    db.commit()
    return {"id": str(ind.id)}


@router.delete("/programs/{prog_id}/indicators/{ind_id}")
def delete_indicator(prog_id: str, ind_id: str, user=Depends(require_supervisor), db: Session = Depends(get_db)):
    ind = db.query(Indicator).filter(Indicator.id == ind_id, Indicator.program_id == prog_id,
                                      Indicator.tenant_id == user["tenant_id"]).first()
    if not ind: raise HTTPException(404, "Indicator not found")
    soft_delete(ind); db.commit()
    return {"deleted": True}


# ── Compute + manual override ───────────────────────────────────────────────

@router.post("/programs/{prog_id}/indicators/{ind_id}/compute")
def compute_indicator_endpoint(prog_id: str, ind_id: str, questionnaire_id: Optional[str] = None,
                                user=Depends(require_supervisor), db: Session = Depends(get_db)):
    """Recompute this indicator's actual values from submissions. If questionnaire_id is
    omitted, recomputes for every wave (questionnaire) under the program."""
    ind = db.query(Indicator).filter(Indicator.id == ind_id, Indicator.program_id == prog_id,
                                      Indicator.tenant_id == user["tenant_id"]).first()
    if not ind: raise HTTPException(404, "Indicator not found")

    q_query = db.query(ProgramQuestionnaire).filter(
        ProgramQuestionnaire.program_id == prog_id, ProgramQuestionnaire.tenant_id == user["tenant_id"])
    if questionnaire_id:
        q_query = q_query.filter(ProgramQuestionnaire.id == questionnaire_id)
    questionnaires = q_query.all()

    written_total = 0
    for q in questionnaires:
        subs = db.query(Submission).filter(
            Submission.questionnaire_id == q.id, Submission.tenant_id == user["tenant_id"]).all()
        written = compute_indicator(ind, q.id, subs, db)
        written_total += len(written)
    db.commit()
    return {"computed": written_total}


@router.put("/programs/{prog_id}/indicator-values/{value_id}")
def override_indicator_value(prog_id: str, value_id: str, body: IndicatorValueOverrideIn,
                              user=Depends(require_supervisor), db: Session = Depends(get_db)):
    val = db.query(IndicatorValue).filter(IndicatorValue.id == value_id,
                                           IndicatorValue.tenant_id == user["tenant_id"]).first()
    if not val: raise HTTPException(404, "Indicator value not found")
    if body.target_value is not None: val.target_value = body.target_value
    if body.actual_value is not None:
        val.actual_value = body.actual_value
        val.value_source = "manual"
    if body.note is not None: val.note = body.note
    db.commit()
    return {"id": str(val.id)}


# ── Indicator Tracking Table (ITT) ──────────────────────────────────────────

def _build_itt(prog_id: str, tenant_id: str, db: Session) -> dict:
    waves = db.query(ProgramQuestionnaire).filter(
        ProgramQuestionnaire.program_id == prog_id, ProgramQuestionnaire.tenant_id == tenant_id,
    ).order_by(ProgramQuestionnaire.wave_number.nullslast(), ProgramQuestionnaire.created_at).all()
    indicators = db.query(Indicator).filter(
        Indicator.program_id == prog_id, Indicator.tenant_id == tenant_id).all()

    wave_out = [{"id": str(w.id), "wave_number": w.wave_number, "wave_label": w.wave_label or w.name} for w in waves]

    ind_rows = []
    for ind in indicators:
        values_by_wave = {}
        for w in waves:
            total_val = db.query(IndicatorValue).filter(
                IndicatorValue.indicator_id == ind.id, IndicatorValue.questionnaire_id == w.id,
                IndicatorValue.disaggregation == {},
            ).first()
            values_by_wave[str(w.id)] = {
                "value_id": str(total_val.id) if total_val else None,
                "target_value": total_val.target_value if total_val else None,
                "actual_value": total_val.actual_value if total_val else None,
                "value_source": total_val.value_source if total_val else None,
            }
        ind_rows.append({
            "id": str(ind.id), "code": ind.code, "name": ind.name, "unit": ind.unit,
            "direction": ind.direction, "logframe_level_id": str(ind.logframe_level_id),
            "values_by_wave": values_by_wave,
        })

    return {"waves": wave_out, "indicators": ind_rows, "logframe": rollup_tree(prog_id, tenant_id, db)}


@router.get("/programs/{prog_id}/itt")
def get_itt(prog_id: str, user=Depends(require_enumerator), db: Session = Depends(get_db)):
    _get_program(prog_id, user, db)
    return _build_itt(prog_id, user["tenant_id"], db)


@router.get("/programs/{prog_id}/itt/xlsx")
def export_itt_xlsx(prog_id: str, user=Depends(require_supervisor), db: Session = Depends(get_db)):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise HTTPException(501, "openpyxl not installed")

    p = _get_program(prog_id, user, db)
    itt = _build_itt(prog_id, user["tenant_id"], db)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Indicator Tracking"

    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    center = Alignment(horizontal="center", vertical="center")

    headers = ["Code", "Indicator", "Unit", "Direction"]
    for w in itt["waves"]:
        label = w["wave_label"] or f"Wave {w['wave_number']}"
        headers += [f"{label} Target", f"{label} Actual"]
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.font = header_font; cell.fill = header_fill; cell.alignment = center

    for ri, ind in enumerate(itt["indicators"], 2):
        vals = [ind["code"], ind["name"], ind["unit"], ind["direction"]]
        for w in itt["waves"]:
            cell = ind["values_by_wave"].get(w["id"], {})
            vals += [cell.get("target_value"), cell.get("actual_value")]
        for ci, v in enumerate(vals, 1):
            ws.cell(row=ri, column=ci, value=v)

    for ci in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(ci)].width = 18

    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    safe = "".join(c for c in p.name if c.isalnum() or c in " _-")
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{safe}_itt.xlsx"'})
